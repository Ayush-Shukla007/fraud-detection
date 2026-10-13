import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re
import os

def markdown_to_docx(md_path, output_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    doc = docx.Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)

    lines = content.split('\n')
    
    in_table = False
    table_data = []
    
    in_code_block = False
    code_data = []
    
    for line in lines:
        if line.startswith('```'):
            if in_code_block:
                in_code_block = False
                p = doc.add_paragraph('\n'.join(code_data))
                p.style.font.name = 'Courier New'
                p.style.font.size = Pt(10)
                code_data = []
            else:
                in_code_block = True
            continue
            
        if in_code_block:
            code_data.append(line)
            continue
            
        if line.startswith('|'):
            in_table = True
            if '---' not in line:
                row = [cell.strip().replace('**', '') for cell in line.split('|') if cell.strip()]
                if row:
                    table_data.append(row)
            continue
        elif in_table:
            # We finished the table
            if table_data:
                table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
                table.style = 'Table Grid'
                for r_idx, row_data in enumerate(table_data):
                    for c_idx, cell_data in enumerate(row_data):
                        table.cell(r_idx, c_idx).text = cell_data
            in_table = False
            table_data = []

        if not line.strip():
            continue
            
        if line.startswith('# '):
            p = doc.add_paragraph(line.replace('# ', '').strip())
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.runs[0].font.size = Pt(16)
            p.runs[0].font.bold = True
        elif line.startswith('## '):
            p = doc.add_paragraph(line.replace('## ', '').strip())
            p.runs[0].font.size = Pt(14)
            p.runs[0].font.bold = True
        elif line.startswith('### '):
            p = doc.add_paragraph(line.replace('### ', '').strip())
            p.runs[0].font.size = Pt(13)
            p.runs[0].font.bold = True
        elif line.startswith('---'):
            pass
        else:
            # Handle bolding via simple regex
            p = doc.add_paragraph()
            parts = re.split(r'(\*\*.*?\*\*)', line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)

    # Change margins to Narrow to match page limits
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    doc.save(output_path)

if __name__ == '__main__':
    md_dir = r'C:\Users\ayush\.gemini\antigravity\brain\23f75955-cec6-4d6a-84c9-62f3d96746c8'
    markdown_to_docx(os.path.join(md_dir, 'FraudGuard_AI_Synopsis.md'), os.path.join(md_dir, 'Updated_Synopsis.docx'))
    markdown_to_docx(os.path.join(md_dir, 'FraudGuard_AI_SRS.md'), os.path.join(md_dir, 'Updated_SRS.docx'))
    print("Successfully generated Word documents.")
